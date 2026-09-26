import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from zipfile import ZipFile
import io

def main():


    # 1. 画面の初期設定
    st.set_page_config(page_title="Excelフォーマット自動修復・結合ツール", page_icon="🛠️", layout="wide")
    st.title("🛠️ Excelフォーマット自動修復・結合ツール")
    st.write("「崩れたファイル」からデータだけを吸い上げ、「正規のファイル」の見た目や数式を活かして完璧なシートを再構築します。")
    
    # 2. VBAマクロが入っているかの判定関数
    def check_vba_inclusion(file):
        if file.name.endswith('.xlsm'):
            return True
        try:
            with ZipFile(file, 'r') as zip_ref:
                if "xl/vbaProject.bin" in zip_ref.namelist():
                    return True
        except Exception:
            pass
        return False
    
    # 3. 2つのファイルをアップロード（画面を2列に分ける）
    col_file1, col_file2 = st.columns(2)
    selected_template_sheet = None
    selected_damaged_sheet = None
    wb_main = None
    wb_data = None
    
    with col_file1:
        template_file = st.file_uploader("🎨 ① 正規のファイル（見本・出力用）をドロップ", type=["xlsx", "xlsm"], key="template")
        # 正規ファイルが置かれたら「すぐに」シート選択を表示
        if template_file:
            has_vba = check_vba_inclusion(template_file)
            wb_main = openpyxl.load_workbook(template_file, data_only=False, keep_vba=has_vba)
            selected_template_sheet = st.selectbox("🎨 見本にする正しいシート", wb_main.sheetnames, index=0)
    
    with col_file2:
        damaged_file = st.file_uploader("❌ ② 崩れたファイル（データ抽出用）をドロップ", type=["xlsx", "xlsm"], key="damaged")
        # 崩れたファイルが置かれたら「すぐに」シート選択を表示
        if damaged_file:
            wb_data = openpyxl.load_workbook(damaged_file, data_only=True, keep_vba=False)
            selected_damaged_sheet = st.selectbox("❌ デザインが崩れたシート", wb_data.sheetnames, index=0)
    
    # 4. 両方のファイルとシートが揃ったら処理用のボタンを表示する
    if template_file and damaged_file and selected_template_sheet and selected_damaged_sheet:
        new_sheet_name = f"{selected_damaged_sheet}_修復済"
        st.info(f"【実行内容】「{selected_template_sheet}」の正規フォーマットに、「{selected_damaged_sheet}」のデータを流し込み、新シート「{new_sheet_name}」を作成します。")
        
        if st.button("✨ 修正済みシートを新規作成する", type="primary"):
            with st.spinner("シートを解析・自動生成中..."):
                source_sheet = wb_main[selected_template_sheet]
                
                if new_sheet_name in wb_main.sheetnames:
                    del wb_main[new_sheet_name]
                    
                ws_target = wb_main.copy_worksheet(source_sheet)
                ws_target.title = new_sheet_name
                has_table = len(source_sheet.tables) > 0
                ws_damaged = wb_data[selected_damaged_sheet]
                
                template_header_map = {}
                for c_idx in range(1, source_sheet.max_column + 1):
                    header_val = source_sheet.cell(row=4, column=c_idx).value or source_sheet.cell(row=1, column=c_idx).value
                    if header_val:
                        template_header_map[str(header_val).strip()] = c_idx
    
                damaged_header_row = None
                damaged_header_map = {}
                max_match_count = 0
    
                # 1〜15行目をスキャンして一番項目名が多く並んでいる行を起点にする
                for r_idx in range(1, min(16, ws_damaged.max_row + 1)):
                    match_count = 0
                    temp_map = {}
                    for c_idx in range(1, ws_damaged.max_column + 1):
                        val = ws_damaged.cell(row=r_idx, column=c_idx).value
                        if val and str(val).strip() in template_header_map:
                            match_count += 1
                            temp_map[c_idx] = str(val).strip()
                    
                    if match_count > max_match_count:
                        max_match_count = match_count
                        damaged_header_row = r_idx
                        damaged_header_map = temp_map
    
                # 見つからない場合の安全弁
                if not damaged_header_row:
                    damaged_header_row = 1
                    for c_idx in range(1, ws_damaged.max_column + 1):
                        val = ws_damaged.cell(row=1, column=c_idx).value
                        if val:
                            damaged_header_map[c_idx] = str(val).strip()
    
                # ----------------------------------------------------
                # データの詰め直し ＆ 枠線自動拡張（分岐処理）
                # ----------------------------------------------------
                template_max_row = source_sheet.max_row
                # データ書き込み開始行（ヘッダー行の次の行からスタート）
                next_target_row = damaged_header_row + 1 
                
                # 壊れたシートのデータ行をループ
                for r_idx in range(damaged_header_row + 1, ws_damaged.max_row + 1):
                    
                    # 空行（勝手に足された中身のない行）は無視する判定
                    row_has_data = False
                    for d_c_idx in list(damaged_header_map.keys()):
                        if ws_damaged.cell(row=r_idx, column=d_c_idx).value is not None:
                            row_has_data = True
                            break
                    if not row_has_data:
                        continue  # 完全な空行ならスキップして上へ詰める
    
                    # 【分岐：テーブルなし（手書き枠）の場合のみPythonが枠を追加】
                    if not has_table and next_target_row > template_max_row:
                        for t_c_idx in template_header_map.values():
                            ref_cell = source_sheet.cell(row=template_max_row, column=t_c_idx)
                            target_cell = ws_target.cell(row=next_target_row, column=t_c_idx)
                            # お手本からスタイル（罫線・背景・フォント等）を完全コピー
                            if ref_cell.has_style:
                                target_cell.font = Font(name=ref_cell.font.name, size=ref_cell.font.size, bold=ref_cell.font.bold, italic=ref_cell.font.italic, color=ref_cell.font.color)
                                target_cell.fill = PatternFill(fill_type=ref_cell.fill.fill_type, start_color=ref_cell.fill.start_color, end_color=ref_cell.fill.end_color)
                                target_cell.border = Border(left=ref_cell.border.left, right=ref_cell.border.right, top=ref_cell.border.top, bottom=ref_cell.border.bottom)
                                target_cell.alignment = Alignment(horizontal=ref_cell.alignment.horizontal, vertical=ref_cell.alignment.vertical)
                                target_cell.number_format = ref_cell.number_format
    
                    # 各列のデータを、起点を元に正しい位置にマッピング
                    for d_c_idx, col_name in damaged_header_map.items():
                        if col_name in template_header_map:
                            target_c_idx = template_header_map[col_name]
                            target_cell = ws_target.cell(row=next_target_row, column=target_c_idx)
                            val = ws_damaged.cell(row=r_idx, column=d_c_idx).value
                            
                            # 【ハイブリッド判定】見本側が数式かつ、値引き行などの手入力データがある場合
                            if target_cell.value and str(target_cell.value).startswith("="):
                                if val is not None:
                                    has_qty_or_price = False
                                    for check_col_name in ["数量", "単価"]:
                                        if check_col_name in template_header_map:
                                            for k, v in damaged_header_map.items():
                                                if v == check_col_name and ws_damaged.cell(row=r_idx, column=k).value is not None:
                                                    has_qty_or_price = True
                                    
                                    if not has_qty_or_price:
                                        # 数量・単価がないのに金額だけある行 ➔ 手入力を優先（値引き等）
                                        target_cell.value = val
                                    else:
                                        # 通常行 ➔ 見本の数式を最優先（手入力は上書きしない）
                                        continue
                                else:
                                    continue
                            else:
                                # 通常のデータ枠ならそのまま値を流し込む
                                target_cell.value = val
                    
                    next_target_row += 1
    
                # 【分岐：テーブルありの場合のみExcelに拡張を任せる（範囲再設定）】
                damaged_max_row = next_target_row - 1
                if has_table:
                    for table_name, table in list(source_sheet.tables.items()):
                        start_cell, end_cell = table.ref.split(":")
                        end_col_letter = "".join([c for c in end_cell if c.isalpha()])
                        new_ref = f"{start_cell}:{end_col_letter}{damaged_max_row}"
                        
                        # コピー元のテーブルと重複しないよう新名称を付けて登録し直す
                        new_table = Table(displayName=f"{table_name}_fixed", ref=new_ref)
                        if table.tableStyleInfo:
                            new_table.tableStyleInfo = table.tableStyleInfo
                        ws_target.add_table(new_table)
    
                # 5. 【修正・完成版】見本の列幅を100%引き継ぎつつ、文字長に合わせて安全に自動拡張
                for col in ws_target.columns:
                    # 列の最初のセル（インデックス0）から列の文字を正しく取得
                    col_letter = col[0].column_letter 
                    
                    # ① 元の見本シートに設定されていたオリジナルの幅を取得する
                    orig_width = source_sheet.column_dimensions[col_letter].width
                    if orig_width is None:
                        orig_width = 11  # 標準の初期幅
    
                    # ② 列全体のセルの文字数（データ長）をスキャンして最大幅を算出する
                    sample_cells = col[:100] if len(col) > 100 else col
                    max_data_len = max(len(str(cell.value or '')) for cell in sample_cells)
                    calculated_width = max_data_len + 3
    
                    # 見本の幅と計算した幅の「広い方」を最終的な列幅として採用する
                    ws_target.column_dimensions[col_letter].width = max(orig_width, calculated_width, 11)
                
                # --- ダウンロード準備 ---
                excel_buffer = io.BytesIO()
                wb_main.save(excel_buffer)
                excel_buffer.seek(0)
                
                download_filename = f"fixed_{template_file.name}"
                mime_type = "application/vnd.ms-excel.sheet.macroEnabled.12" if has_vba else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                
                st.balloons()
                st.success(f"🎉 ファイル内に新シート「{new_sheet_name}」を追加しました！")
                
                if has_table:
                    st.caption("ℹ️ Excelテーブル機能を検出したため、入力のみ行い、範囲設定を自動拡張しました。")
                else:
                    st.caption("ℹ️ 手書き枠を検出したため、Python側で罫線と書式をコピーして拡張しました。")
                
                st.download_button(label="📥 修復済みファイルをダウンロード", data=excel_buffer, file_name=download_filename, mime=mime_type)
    else:
        st.info("💡 「正規のファイル」と「崩れたファイル」の2つをアップロードしてください。")
